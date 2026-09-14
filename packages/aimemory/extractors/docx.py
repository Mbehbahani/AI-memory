"""DOCX extractor via ``python-docx`` (plan section K).

Same failure contract as :mod:`aimemory.extractors.pdf`: never raises, always returns
``ExtractedText(ok=False, reason=...)`` on a corrupt/unreadable file so the source degrades to
``CATALOG_ONLY`` instead of failing the run.
"""

from __future__ import annotations

import io

import docx

from aimemory.domain.ports import ExtractedText

from ._util import finalize_text, truncate_to_bytes

__all__ = ["DocxExtractor"]

_EXTENSIONS = frozenset({".docx"})
_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _iter_table_text(document: "docx.Document") -> list[str]:
    parts: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return parts


class DocxExtractor:
    name = "docx"
    extractor_version = "0.1.0"

    def supports(self, relative_path: str, media_type: str | None = None) -> bool:
        return _suffix(relative_path) in _EXTENSIONS

    def extract(
        self, data: bytes, *, relative_path: str, max_bytes: int | None = None
    ) -> ExtractedText:
        try:
            document = docx.Document(io.BytesIO(data))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            paragraphs.extend(_iter_table_text(document))
            text = "\n\n".join(paragraphs)
        except Exception as exc:  # noqa: BLE001 - python-docx/zipfile raise many distinct types
            return ExtractedText(
                extractor=self.name,
                extractor_version=self.extractor_version,
                media_type=_MEDIA_TYPE,
                ok=False,
                reason=f"docx extraction failed: {type(exc).__name__}",
            )
        if not text.strip():
            return ExtractedText(
                extractor=self.name,
                extractor_version=self.extractor_version,
                media_type=_MEDIA_TYPE,
                ok=False,
                reason="docx produced no extractable text (empty document)",
            )
        truncated = False
        if max_bytes is not None:
            text, truncated = truncate_to_bytes(text, max_bytes)
        stored, char_count = finalize_text(text)
        return ExtractedText(
            text=stored,
            extractor=self.name,
            extractor_version=self.extractor_version,
            media_type=_MEDIA_TYPE,
            char_count=char_count,
            truncated=truncated,
            ok=True,
        )


def _suffix(relative_path: str) -> str:
    name = relative_path.replace("\\", "/").rsplit("/", 1)[-1]
    idx = name.rfind(".")
    return name[idx:].lower() if idx != -1 else ""
